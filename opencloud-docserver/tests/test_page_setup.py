"""Page setup round-trip tests (F-090 page size / F-091 orientation / F-092 margins).

The .page-setup marker div at body start carries document page geometry
(twips) and maps to DOCX w:sectPr (w:pgSz + w:pgMar) and ODT
style:page-layout-properties. Marker absence = unchanged defaults.
"""

import io

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.shared import Inches

from src.editor.converter import docx_to_html, html_to_docx
from src.editor.odt_converter import html_to_odt, odt_to_html
from src.editor.sanitize import sanitize_html

MARK_A4L = (
    '<div class="page-setup" data-page-w="16838" data-page-h="11906" '
    'data-orient="landscape" data-margin-top="1134" data-margin-bottom="1134" '
    'data-margin-left="1417" data-margin-right="1417"></div>'
)


def test_docx_emits_page_setup_marker():
    d = Document()
    s = d.sections[0]
    s.page_width = Inches(11.69)
    s.page_height = Inches(8.27)
    s.orientation = WD_ORIENT.LANDSCAPE
    s.top_margin = Inches(0.8)
    s.left_margin = Inches(1.2)
    d.add_paragraph("hello")
    buf = io.BytesIO()
    d.save(buf)
    html = docx_to_html(buf.getvalue())
    assert html.startswith('<div class="page-setup"')
    assert 'data-orient="landscape"' in html
    assert 'data-margin-top="1152"' in html  # 0.8in
    assert 'data-margin-left="1728"' in html  # 1.2in


def test_docx_marker_applies_to_section_and_is_stripped_from_body():
    doc = Document(io.BytesIO(html_to_docx(MARK_A4L + "<p>body</p>")))
    s = doc.sections[0]
    assert abs(s.page_width.inches - 11.69) < 0.01
    assert abs(s.page_height.inches - 8.27) < 0.01
    assert s.orientation == WD_ORIENT.LANDSCAPE
    assert abs(s.top_margin.inches - 1134 / 1440) < 0.01
    assert [p.text for p in doc.paragraphs] == ["body"]  # marker not a paragraph


def test_docx_without_marker_keeps_defaults():
    doc = Document(io.BytesIO(html_to_docx("<p>plain</p>")))
    assert doc.sections[0].page_width == Inches(8.5)  # python-docx Letter default
    assert doc.sections[0].orientation != WD_ORIENT.LANDSCAPE


def test_odt_page_setup_roundtrip_exact():
    html = odt_to_html(html_to_odt(MARK_A4L + "<p>Hi odt</p>"))
    assert 'data-page-w="16838"' in html
    assert 'data-page-h="11906"' in html
    assert 'data-orient="landscape"' in html
    assert 'data-margin-left="1417"' in html
    assert "<p>Hi odt</p>" in html


def test_page_setup_marker_survives_sanitize():
    clean = sanitize_html(MARK_A4L + "<p>x</p>")
    assert 'class="page-setup"' in clean
    assert 'data-page-w="16838"' in clean


def test_page_setup_coexists_with_header_footer():
    html = MARK_A4L + '<header class="page-header"><p>H</p></header><p>b</p>'
    doc = Document(io.BytesIO(html_to_docx(html)))
    assert abs(doc.sections[0].page_width.inches - 11.69) < 0.01
    assert doc.sections[0].header.paragraphs[0].text == "H"
    oh = odt_to_html(html_to_odt(html))
    assert 'data-page-w="16838"' in oh and ">H<" in oh


def test_docx_default_geometry_emits_no_marker():
    """Marker absence = defaults; emitting one for every default document
    would churn every exact-HTML assert in the corpus."""
    d = Document()
    d.add_paragraph("plain")
    buf = io.BytesIO()
    d.save(buf)
    assert not docx_to_html(buf.getvalue()).startswith('<div class="page-setup"')
