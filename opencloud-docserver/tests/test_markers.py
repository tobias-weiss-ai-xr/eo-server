"""Section-flag + drop cap + borders round-trip tests.

Marker contracts (all ride at body start, in this fixed order, and are
stripped from the rendered body):
  * hyphenation:  <div class="hyphenation" data-auto="1">            -> w:autoHyphenation (settings.xml)
  * line numbers: <div class="line-numbers" data-restart data-distance> -> w:lnNumType (sectPr)
  * watermark:    <div class="watermark" data-text data-color>     -> centered 32pt bold gray header paragraph
Marker absence = the feature is off; marker-less documents stay byte-stable.
Drop cap rides inside the paragraph: <span class="dropcap">X</span> as the
first run -> w:framePr w:dropCap. Paragraph borders map to w:pBdr / ODT
fo:border-*.
"""

import io

from docx import Document

from src.editor.converter import docx_to_html, html_to_docx
from src.editor.odt_converter import html_to_odt, odt_to_html
from src.editor.sanitize import sanitize_html


def _rt(html: str) -> str:
    return docx_to_html(html_to_docx(sanitize_html(html)))


# ── hyphenation ────────────────────────────────────────────────────────
def test_hyphenation_marker_round_trip():
    h = '<div class="hyphenation" data-auto="1"></div><p>Hyphenated.</p>'
    assert _rt(h).split("\n")[0] == '<div class="hyphenation" data-auto="1"></div>'


def test_hyphenation_absent_when_off():
    out = _rt("<p>Plain.</p>")
    assert "hyphenation" not in out
    assert out == "<p>Plain.</p>"


def test_hyphenation_written_to_settings_xml():
    data = html_to_docx('<div class="hyphenation" data-auto="1"></div><p>x</p>')
    d = Document(io.BytesIO(data))
    assert d.settings.element.find(
        '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
        'autoHyphenation') is not None


# ── line numbers ───────────────────────────────────────────────────────
def test_line_numbers_marker_round_trip():
    h = ('<div class="line-numbers" data-restart="eachPage" data-distance="720">'
         '</div><p>Numbered.</p>')
    out = _rt(h).split("\n")[0]
    assert out == ('<div class="line-numbers" data-restart="eachPage" '
                   'data-distance="720"></div>')


def test_line_numbers_absent_when_off():
    assert "line-numbers" not in _rt("<p>No numbers.</p>")


def test_line_numbers_written_to_sectpr():
    data = html_to_docx(
        '<div class="line-numbers" data-restart="continuous" data-distance="480">'
        '</div><p>x</p>')
    d = Document(io.BytesIO(data))
    ln = d.sections[0]._sectPr.find(
        '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
        'lnNumType')
    assert ln is not None
    assert ln.get(
        '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
        'restart') == "continuous"
    assert ln.get(
        '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
        'distance') == "480"


# ── watermark ──────────────────────────────────────────────────────────
def test_watermark_marker_round_trip_docx():
    h = ('<div class="watermark" data-text="CONFIDENTIAL" data-color="#C0C0C0">'
         '</div><p>Body.</p>')
    out = _rt(h)
    lines = out.split("\n")
    assert lines[0] == ('<div class="watermark" data-text="CONFIDENTIAL" '
                        'data-color="#c0c0c0"></div>')
    assert "CONFIDENTIAL" not in "<p>Body.</p>" + "\n".join(lines[1:]).replace(
        "CONFIDENTIAL", "")


def test_watermark_coexists_with_header():
    h = ('<div class="watermark" data-text="TOP SECRET" data-color="#999999">'
         '</div><header><p>My header</p></header><p>Body.</p>')
    out = _rt(h)
    lines = out.split("\n")
    assert lines[0].startswith('<div class="watermark"')
    assert '<p>My header</p>' in "\n".join(lines[1:])


def test_watermark_absent_without_marker():
    assert "watermark" not in _rt("<p>No marker here.</p>")


def test_watermark_round_trip_odt():
    h = ('<div class="watermark" data-text="SECRET" data-color="#C0C0C0">'
         '</div><p>ODT body.</p>')
    out = odt_to_html(html_to_odt(sanitize_html(h)))
    # ODT keeps the text; the color is an ODT divergence (documented L1).
    assert out.split("\n")[0] == '<div class="watermark" data-text="SECRET"></div>'
    assert "SECRET" not in out.split("\n", 1)[1]


def test_watermark_odt_absent_without_marker():
    assert "watermark" not in odt_to_html(html_to_odt("<p>plain</p>"))


# ── drop cap ───────────────────────────────────────────────────────────
def test_dropcap_round_trip_docx():
    h = '<p><span class="dropcap">H</span>ello world.</p>'
    assert _rt(h) == '<p><span class="dropcap">H</span>ello world.</p>'


def test_dropcap_absent_without_marker():
    assert "dropcap" not in _rt("<p>Normal.</p>")


def test_dropcap_letters_kept_on_write():
    data = html_to_docx('<p><span class="dropcap">W</span>ord.</p>')
    d = Document(io.BytesIO(data))
    assert d.paragraphs[0].text == "Word."
    fp = d.paragraphs[0]._p.pPr.find(
        '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
        'framePr')
    assert fp is not None
    assert fp.get(
        '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
        'dropCap') == "drop"


# ── paragraph borders ──────────────────────────────────────────────────
def test_border_shorthand_round_trip_docx():
    h = '<p style="border:1.25pt solid #ff0000">Boxed.</p>'
    out = _rt(h)
    for side in ("top", "left", "bottom", "right"):
        assert f"border-{side}:1.25pt solid #ff0000" in out


def test_border_sides_round_trip_docx():
    h = ('<p style="border-top:2pt solid #0000ff;border-bottom:2pt solid '
         '#0000ff">TB.</p>')
    out = _rt(h)
    assert "border-top:2pt solid #0000ff" in out
    assert "border-bottom:2pt solid #0000ff" in out
    assert "border-left" not in out


def test_border_absent_without_style():
    assert "border" not in _rt("<p>No frame here.</p>")


def test_border_round_trip_odt():
    h = '<p style="border:1.5pt solid #0000ff">Boxed ODT.</p>'
    out = odt_to_html(html_to_odt(sanitize_html(h)))
    assert "border-top:1.5pt solid #0000ff" in out
    assert "border-bottom:1.5pt solid #0000ff" in out


# ── marker stripping guarantees ────────────────────────────────────────
def test_markers_never_leak_into_body():
    for h in (
        '<div class="hyphenation" data-auto="1"></div><p>x</p>',
        '<div class="line-numbers" data-restart="eachPage"></div><p>x</p>',
        '<div class="watermark" data-text="T" data-color="#C0C0C0"></div><p>x</p>',
        '<p style="border:1pt solid #000">x</p>',
        '<p><span class="dropcap">X</span>y</p>',
    ):
        body = docx_to_html(html_to_docx(sanitize_html(h)))
        assert "\n" not in body or "<div class=" in body.split("\n")[0]
    # the marker divs themselves are not in the body paragraph stream
    out = docx_to_html(html_to_docx(
        '<div class="hyphenation" data-auto="1"></div>'
        '<div class="line-numbers"></div>'
        '<div class="watermark" data-text="T" data-color="#C0C0C0"></div>'
        '<p>real content</p>'))
    assert out.count("<div class=") == 3
    assert "<p>real content</p>" in out
